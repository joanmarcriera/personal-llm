from __future__ import annotations

import csv
import json
import mailbox
from email import policy
from email.parser import BytesParser
from pathlib import Path
from typing import Any

import pandas as pd
from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from markdownify import markdownify
from pypdf import PdfReader

from personal_llm.connectors.email import EmailConnector
from personal_llm.core.io import read_text, sha256_text
from personal_llm.core.schemas import ExtractedDocument, SourceDocument, SourceType


class DocumentExtractor:
    def extract(self, source: SourceDocument) -> list[ExtractedDocument]:
        path = Path(source.content_path)
        if source.source_type == SourceType.LINKWARDEN_EXPORT:
            return self._extract_linkwarden(source, path)
        if path.suffix.lower() == ".mbox":
            return self._extract_mbox(source, path)
        if path.suffix.lower() == ".eml":
            return [self._extract_eml(source, path.read_bytes())]
        if path.suffix.lower() == ".pdf":
            return [self._single(source, self._extract_pdf(path))]
        if path.suffix.lower() == ".docx":
            return [self._single(source, self._extract_docx(path))]
        if path.suffix.lower() in {".html", ".htm"}:
            return [self._single(source, self._extract_html(path))]
        if path.suffix.lower() in {".xlsx", ".csv"}:
            return [self._single(source, self._extract_spreadsheet(path))]
        if path.suffix.lower() in {".md", ".markdown"}:
            return [self._single(source, read_text(path))]
        return [self._single(source, read_text(path))]

    def _single(self, source: SourceDocument, text: str) -> ExtractedDocument:
        return ExtractedDocument(
            id=f"doc::{source.id}",
            source_id=source.id,
            source_type=source.source_type,
            title=source.title,
            text=text.strip(),
            metadata=source.metadata,
            checksum=sha256_text(text),
        )

    def _extract_pdf(self, path: Path) -> str:
        reader = PdfReader(str(path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    def _extract_docx(self, path: Path) -> str:
        document = DocxDocument(path)
        return "\n".join(paragraph.text for paragraph in document.paragraphs)

    def _extract_html(self, path: Path) -> str:
        soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="ignore"), "html.parser")
        return markdownify(str(soup))

    def _extract_spreadsheet(self, path: Path) -> str:
        if path.suffix.lower() == ".csv":
            with path.open("r", encoding="utf-8", errors="ignore") as handle:
                return "\n".join(",".join(row) for row in csv.reader(handle))
        workbook = pd.read_excel(path, sheet_name=None)
        lines: list[str] = []
        for sheet_name, frame in workbook.items():
            lines.append(f"# Sheet: {sheet_name}")
            lines.append(frame.fillna("").to_csv(index=False))
        return "\n".join(lines)

    def _extract_linkwarden(self, source: SourceDocument, path: Path) -> list[ExtractedDocument]:
        payload = json.loads(path.read_text(encoding="utf-8"))
        records = payload.get("links", []) if isinstance(payload, dict) else []
        documents: list[ExtractedDocument] = []
        for index, record in enumerate(records):
            if not isinstance(record, dict):
                continue
            title = str(record.get("name") or record.get("url") or f"bookmark-{index}")
            text = "\n".join(
                [
                    title,
                    str(record.get("description") or ""),
                    str(record.get("url") or ""),
                    str(record.get("textContent") or record.get("readableContent") or ""),
                ]
            ).strip()
            metadata = {
                "url": record.get("url"),
                "tags": record.get("tags", []),
                "collection": record.get("collection"),
                "archived": bool(record.get("archived")),
            }
            documents.append(
                ExtractedDocument(
                    id=f"{source.id}::bookmark::{index}",
                    source_id=source.id,
                    source_type=source.source_type,
                    title=title,
                    text=text,
                    metadata=metadata,
                    checksum=sha256_text(text),
                )
            )
        return documents

    def _extract_mbox(self, source: SourceDocument, path: Path) -> list[ExtractedDocument]:
        box = mailbox.mbox(path)
        documents: list[ExtractedDocument] = []
        for index, message in enumerate(box):
            payload = self._render_email(message.as_bytes())
            summary = EmailConnector.summarize_message(payload["message"])
            documents.append(
                ExtractedDocument(
                    id=f"{source.id}::message::{index}",
                    source_id=source.id,
                    source_type=source.source_type,
                    title=str(summary.get("subject") or f"message-{index}"),
                    text=payload["text"],
                    metadata=summary | {"attachments": payload["attachments"]},
                    checksum=sha256_text(payload["text"]),
                )
            )
        return documents

    def _extract_eml(self, source: SourceDocument, raw_message: bytes) -> ExtractedDocument:
        payload = self._render_email(raw_message)
        summary = EmailConnector.summarize_message(payload["message"])
        text = payload["text"]
        return ExtractedDocument(
            id=f"{source.id}::eml",
            source_id=source.id,
            source_type=source.source_type,
            title=str(summary.get("subject") or source.title),
            text=text,
            metadata=summary | {"attachments": payload["attachments"]},
            checksum=sha256_text(text),
        )

    def _render_email(self, raw_message: bytes) -> dict[str, Any]:
        message = BytesParser(policy=policy.default).parsebytes(raw_message)
        body_parts: list[str] = []
        attachments: list[dict[str, Any]] = []
        for part in message.walk():
            content_disposition = part.get_content_disposition()
            content_type = part.get_content_type()
            if content_disposition == "attachment":
                attachments.append(
                    {
                        "filename": part.get_filename(),
                        "content_type": content_type,
                        "size": len(part.get_payload(decode=True) or b""),
                    }
                )
                continue
            if content_type == "text/plain":
                body_parts.append(part.get_content())
            elif content_type == "text/html":
                body_parts.append(markdownify(part.get_content()))
        if not body_parts:
            body_parts.append(
                message.get_body(preferencelist=("plain", "html")).get_content()
                if message.get_body()
                else ""
            )
        text = "\n".join(part for part in body_parts if part).strip()
        return {"message": message, "text": text, "attachments": attachments}
